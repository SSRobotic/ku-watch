"""Minimal evidence-backed revision; retain the earlier checklist unchanged."""
from pathlib import Path
from datetime import datetime,timezone,timedelta
import json,hashlib
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'output/deadline_20261018/requirements_checklist_20261014'
OUT=ROOT/'output/deadline_20261018/requirements_accepted_20261014'
def fmt(ms):
 minutes=int(ms//60000);return f'{minutes//60} h {minutes%60} min'
def set_text(p,text,size=11):
 p.text=text;p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(0)
 for r in p.runs:r.font.size=Pt(size)
def main():
 OUT.mkdir(parents=True,exist_ok=True)
 frozen=OUT/'frozen_sources.json'
 if not frozen.exists():
  files=['output/accumulated_soak/status.json','output/deadline_20261018/fr02_visual_acceptance.json','output/deadline_20261018/hub_acceptance_now.json','output/deadline_20261018/sd_acceptance_longer/env_1791956522_e47806b1.csv.audit.json','output/deadline_20261018/sd_acceptance_longer/excel_open.json','output/deadline_20261018/observed_boot_change_20261014.json']
  data={'frozen_utc':datetime.now(timezone.utc).isoformat(),'snapshots':{p:json.loads((ROOT/p).read_text(encoding='utf-8-sig')) for p in files},'hashes':{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in files}}
  frozen.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
 data=json.loads(frozen.read_text(encoding='utf8'));snap=data['snapshots']
 soak=snap['output/accumulated_soak/status.json'];sd=snap['output/deadline_20261018/sd_acceptance_longer/env_1791956522_e47806b1.csv.audit.json']
 assert snap['output/deadline_20261018/fr02_visual_acceptance.json']['fr02_functional_acceptance_pass']
 assert snap['output/deadline_20261018/sd_acceptance_longer/excel_open.json']['passed'] and sd['five_second_cadence'] and sd['all_columns_valid']
 when=datetime.fromisoformat(soak['updated_utc'].replace('Z','+00:00')).astimezone(timezone(timedelta(hours=7)))
 original=SOURCE/'KU_Watch_ตารางเช็กRequirements_14ตุลาคม2569.docx';before=hashlib.sha256(original.read_bytes()).hexdigest();doc=Document(original)
 rows={r.cells[0].text.splitlines()[0]:r for t in doc.tables for r in t.rows[1:]}
 acceptance=rows['FR-02'].cells[3]
 checked=acceptance._tc.xpath('.//w14:checkbox/w14:checked')[0];checked.set(qn('w14:val'),'1')
 acceptance._tc.xpath('.//w:sdtContent/w:r/w:t')[0].text='☑';set_text(acceptance.paragraphs[1],'ผ่าน',10)
 set_text(rows['FR-02'].cells[4].paragraphs[0],'ผู้ใช้ยืนยันเวลาปัจจุบัน/แบตเป็นตัวเลข และ Wi-Fi/MQTT ตรงจริง รวมคำยืนยันทัช\nHUB 1.0.7 มีภาพ/รับข้อมูล/ACK ปกติ; ผ่านการใช้งาน FR-02 แต่ยังไม่ใช่ผลทัช ≤300 ms ของ NFR-03')
 set_text(rows['FR-04'].cells[3].paragraphs[1],'ช่วงที่ตรวจ',10)
 seconds=sd['duration_s'];span=f'{seconds//3600} h {seconds%3600//60} min {seconds%60} s'
 set_text(rows['FR-04'].cells[4].paragraphs[0],f"ไฟล์การ์ดจริง {sd['rows']:,} แถว/{span}; Excel 26 คอลัมน์; เวลา RTC ต่าง ≤{sd['rtc_watch_max_difference_s']} s\nรอบบันทึก ~5 s; ผลต่างเวลา RTC จำนวนเต็ม 4–6 s ไม่มีช่วงขาด >6 s; ผลจำกัดช่วงที่ตรวจ ไม่แทน 24 h/แบต")
 set_text(rows['NFR-01'].cells[4].paragraphs[0],f"MQTT สะสม {fmt(soak['credited_ms'])}; ต่อเนื่องยาวที่สุด {fmt(soak['longest_continuous_ms'])}; ช่วงปัจจุบัน {fmt(soak['current_continuous_ms'])}\nพบรหัสรอบบูตเปลี่ยน 14:41 น. และช่วงรับห่าง 34.928 s จึงเริ่มช่วงต่อเนื่องใหม่ ยังไม่ผ่าน 24 h")
 for p in doc.paragraphs:
  if p.text.startswith('สถานะ 14 ตุลาคม'):set_text(p,f'สถานะ 14 ตุลาคม 2569 เวลา {when:%H:%M} น.   WATCH 2.0.8   HUB 1.0.7   ฉบับปรับผลตรวจรับ',10.5)
 for s in doc.sections:
  for p in s.footer.paragraphs:
   for r in p.runs:
    if '14 ตุลาคม 2569' in r.text:r.text=f'14 ตุลาคม 2569 {when:%H:%M} น.   หน้า ';r.font.size=Pt(9)
 # Clarify this is the project's acceptance plan, not a new instructor mandate.
 last=doc.paragraphs[-1]
 set_text(last,'ต้นฉบับ Requirements V2 ไม่เปลี่ยนเกณฑ์ ตารางนี้เทียบเป้าหมายโครงการ ไม่ยืนยันว่าผู้สอนบังคับทุกข้อในงาน Proposal ข้อที่พัก/ยังไม่ทดสอบไม่ติ๊กผ่าน Garmin รอชาร์จ อุณหภูมิยังไม่มีเครื่องอ้างอิง และราคาจริงยังต้องเติม',10.5)
 file=OUT/'KU_Watch_ตารางเช็กRequirements_ตรวจรับล่าสุด_14ตุลาคม2569.docx';doc.save(file)
 assert hashlib.sha256(original.read_bytes()).hexdigest()==before
 compare=json.loads((SOURCE/'comparison.json').read_text(encoding='utf8'))
 for r in compare['requirements']:
  if r['id']=='FR-02':r['acceptance_checked']=True;r['acceptance_label']='ผ่าน';r['evidence_and_remaining']=rows['FR-02'].cells[4].text
  if r['id']=='FR-04':r['acceptance_label']='ช่วงที่ตรวจ';r['evidence_and_remaining']=rows['FR-04'].cells[4].text
  if r['id']=='NFR-01':r['evidence_and_remaining']=rows['NFR-01'].cells[4].text
 compare.update(snapshot_utc=soak['updated_utc'],snapshot_thai=when.isoformat(),boundary='FR-02 functional visual acceptance and FR-04 actual CSV observation checked. Other partial/deferred criteria remain unchecked. Project acceptance targets are separate from instructor proposal grading.')
 (OUT/'comparison.json').write_text(json.dumps(compare,ensure_ascii=False,indent=2),encoding='utf8')
 (OUT/'manifest.json').write_text(json.dumps({'docx':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'original_retained_sha256':before,'checked_acceptance':['FR-02','FR-04'],'source_evidence':str(frozen)},indent=2),encoding='utf8')
 print(file)
if __name__=='__main__':main()
