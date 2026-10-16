"""Convert the frozen teammate progress brief into an editable Word document."""
from pathlib import Path
import re,json,hashlib
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
FOLDER=ROOT/'output/deadline_20261018/slide_progress_20261014T065939Z'
SOURCE=FOLDER/'ความคืบหน้า_ส่งเพื่อนทำสไลด์.md'
OUT=FOLDER/'KU_Watch_ความคืบหน้าสำหรับทำสไลด์_14ตุลาคม2569.docx'
FONT='Leelawadee UI'
def inline(p,text):
    for i,part in enumerate(re.split(r'\*\*',text)):
        r=p.add_run(part);r.bold=bool(i%2)
def cell_property(cell,name,attrs):
    prop=OxmlElement('w:'+name)
    for k,v in attrs.items():prop.set(qn('w:'+k),str(v))
    cell._tc.get_or_add_tcPr().append(prop)
def table(doc,lines):
    data=[[x.strip() for x in line.strip('|').split('|')] for line in lines]
    data=[row for row in data if not all(re.fullmatch(r'[-:]+',s) for s in row)]
    t=doc.add_table(rows=0,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=[Inches(1.35),Inches(5.65)]
    for col,width in zip(t.columns,widths):col.width=width
    for idx,rowdata in enumerate(data):
        row=t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if idx==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(row.cells,rowdata)):
            cell.width=widths[j];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_property(cell,'shd',{'fill':'DCE6EF' if idx==0 else 'FFFFFF'})
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                item=OxmlElement('w:'+side);item.set(qn('w:w'),'100');item.set(qn('w:type'),'dxa');margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side)
                for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items():b.set(qn('w:'+k),v)
                borders.append(b)
            cell._tc.get_or_add_tcPr().append(borders)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            inline(p,text)
            for r in p.runs:r.font.size=Pt(11);r.bold=idx==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def main():
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.75)
    section.header_distance=section.footer_distance=Inches(.3)
    for style in ['Normal','Title','Heading 1','Heading 2','List Bullet','List Number']:
        s=doc.styles[style];s.font.name=FONT;s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
        s.element.get_or_add_rPr().rFonts.set(qn('w:cs'),FONT)
    normal=doc.styles['Normal'];normal.font.size=Pt(11.5)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
    for name in ['List Bullet','List Number']:
        doc.styles[name].font.size=Pt(11.5);doc.styles[name].paragraph_format.space_after=Pt(6)
        doc.styles[name].paragraph_format.line_spacing=1.12
    doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(10)
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
    header=section.header.paragraphs[0];header.text='KU Watch   ข้อมูลสำหรับเตรียมสไลด์';header.runs[0].font.size=Pt(9)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('14 ตุลาคม 2569   หน้า ');r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='ความคืบหน้า KU Watch สำหรับทำสไลด์'
    doc.core_properties.author='KU Watch'
    lines=SOURCE.read_text(encoding='utf8').splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            table(doc,block);continue
        if line.startswith('# '):
            doc.add_paragraph('ความคืบหน้า KU Watch สำหรับทำสไลด์',style='Title');continue
        if line.startswith('## '):
            heading=line[3:]
            if heading in ['สถานะการทดสอบที่ใส่สไลด์ได้','งานค้างก่อนสรุปว่าครบ Requirements']:doc.add_page_break()
            doc.add_paragraph(heading,style='Heading 1');continue
        if line.startswith('- '):p=doc.add_paragraph(style='List Bullet');inline(p,line[2:]);continue
        if re.match(r'^\d+\. ',line):p=doc.add_paragraph(style='List Number');inline(p,re.sub(r'^\d+\. ','',line));continue
        if line.startswith('อัปเดต 14 ตุลาคม'):
            line='อัปเดต 14 ตุลาคม 2569 เวลา 13:59 น. ตัวเลขเป็นสถานะ ณ เวลาที่ระบุ และอาจเปลี่ยนภายหลัง'
        p=doc.add_paragraph();inline(p,line)
    for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.save(OUT)
    manifest={'created_utc':datetime.now(timezone.utc).isoformat(),'source':SOURCE.name,'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'docx':OUT.name,'docx_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'snapshot_preserved':'14 Oct 2026 13:59 Asia/Bangkok; document conversion does not refresh or change evidence.'}
    (FOLDER/'word_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')
    print(OUT)
if __name__=='__main__':main()
