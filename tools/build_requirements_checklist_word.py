"""Build an editable, evidence-scoped checklist against the original 16 criteria."""
from pathlib import Path
from datetime import datetime,timezone,timedelta
import json,hashlib,argparse
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/deadline_20261018/requirements_checklist_20261014'
FONT='Leelawadee UI';W14='http://schemas.microsoft.com/office/word/2010/wordml'
ROWS=[
 ('FR-01','Must','MAX30102 ชีพจรและ SpO₂; QMI8658 ความเร่งและไจโร; BME680 อุณหภูมิ/ความชื้น\nเกณฑ์: ชีพจร/SpO₂/การเคลื่อนไหว ≥1 Hz, IAQ ทุก 3 s และทำงานครบ 1 นาที',False,'บางส่วน',False,'ยังไม่ครบ','มีข้อมูลสิ่งแวดล้อม/การเคลื่อนไหว/INA/IAQ สด; MAX30102 มีสัญญาณดิบ แต่รอบล่าสุดไม่มี BPM/SpO₂ ที่ valid ต้องเก็บช่วงอ่านได้จริง\nใช้ GY-51 ตามการปรับที่ผู้ใช้ยืนยัน ไม่มีไจโรจาก GY-51'),
 ('FR-02','Must','แสดงค่าล่าสุด เวลา แบต INA226 และสถานะ Wi-Fi บนจอทัช\nเกณฑ์: รีเฟรชทุก 500 ms แสดงสถานะตรงจริงและแตะสั่งงานได้',True,'แล้ว',False,'รอตรวจครบ','ผู้ใช้ยืนยันภาพ เวลา และทัช HUB 1.0.7; ESP-NOW ส่งคำสั่งเปลี่ยนหน้ามี ACK\nยังต้องเก็บชุดตรวจค่าล่าสุด/แบต/ไอคอน Wi-Fi ครบในรุ่นปัจจุบัน แยกจากการจับเวลาทัช'),
 ('FR-03','Must','JSON ผ่าน MQTT ทุก 1 s, QoS 1, LWT /status และ topic\nku/iot/<group>/\n<device>/env\nเกณฑ์: รับครบ 5 นาที',False,'บางส่วน',False,'รอตรวจ','MQTT สดและส่งย้อนหลังทำงาน; ปัจจุบันใช้ ku_watch/live ซึ่งยังต่างจากรูปแบบ topic เดิม\nปรับ topic และฝั่งรับหลังรอบยาว แล้วตรวจ 5 นาทีตามชื่อใหม่'),
 ('FR-04','Must','MicroSD CSV ทุก 5 s อ้างเวลา PCF85063 และปิดไฟล์เป็นระยะ\nเกณฑ์: เปิด Excel ได้ เวลาถูกต้องและไม่มีแถวขาด',True,'แล้ว',True,'ช่วงสั้น','อ่านไฟล์จริงจากการ์ด 18 แถว/86 s; Excel ครบ 26 คอลัมน์; RTC ต่าง ≤1 s; errors/dropped=0\nรอบแรก 6 s จากนั้น 5 s ทุกแถว ผลผ่านนี้จำกัดช่วงสั้น ไม่แทนผลระยะยาว'),
 ('FR-05','Should','เมื่อ Wi-Fi/broker หลุด เก็บ MicroSD/NVS และส่งข้อมูลค้างกลับ MQTT อัตโนมัติ\nเกณฑ์: ตัด Wi-Fi 10 นาที ข้อมูลครบและเซ็นเซอร์ไม่หยุด',True,'แล้ว',False,'รอตรวจ','มีแฟลชคิวตัวเล็กและผลส่งย้อนหลังจริง ชุดที่สังเกต 561 รหัสเข้า DB ครบอย่างละหนึ่งแถว\nยังต้องตัด 10 นาทีในรุ่นปัจจุบันและตรวจรหัสทั้งช่วง; ไม่ส่งแฟลชคิวย้อนเข้าการ์ดเมื่อ hub ปิด'),
 ('FR-06','Should','IAQ >150 หรืออุณหภูมิ >35°C ต่อเนื่อง 3 รอบอ่านใหม่ เตือนบนจอและ dashboard\nเกณฑ์: ป้อนค่าจำลองแล้วเตือนภายใน 5 s',True,'แล้ว',False,'รอตรวจ','model/renderer ผ่าน fixture ป้อนค่าใหม่; ต้องจับเวลาถึงจอจริงและ Grafana\nBSEC จริงอ่านราว 3 s/รอบ การรอ 3 รอบอาจเกิน 5 s ต้องแยกผลจำลองจากผลอ่านจริง'),
 ('FR-07','Could','ตั้ง SSID/password, broker และเกณฑ์เตือนโดยไม่แฟลช ใช้ NVS/Preferences\nเกณฑ์: เปลี่ยนแล้วรีบูต ค่ายังอยู่และมีผลทันที',True,'แล้ว',False,'รอตรวจ','มีเว็บ SETUP/NVS และเคยเปลี่ยนเครือข่ายใช้งานได้\nยังต้องตรวจการคงค่าทั้งชุดหลังเปิดใหม่ในรุ่นเดียวกัน ไม่รีบูตตัวเล็กกลางรอบยาว'),
 ('FR-08','Could','เช็ก OTA ตอนบูตและทุก 24 h ดาวน์โหลดรุ่นใหม่เอง แสดงเวอร์ชันและรายงาน /status\nเกณฑ์: อัปเดตรุ่นใหม่เองภายใน 24 h',False,'บางส่วน',False,'รอตรวจ','มี OTA service/manifest/ไฟล์เฟิร์มแวร์และเลขเวอร์ชัน\nยังไม่มีผลดาวน์โหลดรุ่นใหม่เองครบวงจร ต้องมี HTTPS ที่ใช้งานได้; แม้เป็น Could ยังสัมพันธ์กับ NFR-08'),
 ('NFR-01','คุณภาพ','เดินระบบต่อเนื่องอย่างน้อย 24 h ไม่ต้องรีบูตด้วยมือ\nเกณฑ์: ไม่ค้างและยังส่งข้อมูลได้ปกติ',False,'ทดสอบอยู่',False,'รอครบ','SOAK_DYNAMIC'),
 ('NFR-02','คุณภาพ','ออฟไลน์ยังทำงาน ต่อกลับด้วย exponential backoff\nเกณฑ์: ปิด broker/เราเตอร์ 10 นาที ไม่ค้างและส่งกลับเอง',True,'แล้ว',False,'รอตรวจ','มี backoff และต่อกลับเองจริง\nต้องผูกผลตัด 10 นาทีและข้อมูลย้อนหลังเข้ารุ่นปัจจุบัน ใช้รอบเดียวกับ FR-05 ได้'),
 ('NFR-03','คุณภาพ','เซ็นเซอร์ถึงภาพ Grafana ≤5 s และทัช ≤300 ms\nเกณฑ์: วัด Grafana 10 ครั้ง เฉลี่ย ≤5 s แตะไม่หน่วง',True,'แล้ว',False,'รอวัดครบ','ผล query เดิม 10 รหัสเฉลี่ย 3.091 s ยังไม่รวมภาพ browser\nทัชอ่านทุก ~4 ms/อายุเหตุการณ์สูงสุด 103 ms เป็นค่าภายใน ต้องวัดแตะถึงภาพจริงและ browser 10 ครั้ง'),
 ('NFR-04','คุณภาพ','LiPo 1000 mAh ใช้ได้ ≥6 h พร้อมหรี่/พักจอ และบันทึก/ส่งข้อมูล\nเกณฑ์: จับเวลาจนแบตเหลือ 20% เทียบแรงดัน INA226',False,'พักไว้',False,'ไม่ทดสอบ','ผู้ใช้ไม่มีที่ชาร์จ/เวลาไม่พอ รอบนี้ใช้ USB; นอกสถานที่ใช้แบต 1000 mAh และเปลี่ยนเมื่อหมด\nUSB ไม่เป็นผลอายุแบต 6 h ไม่เปลี่ยนเกณฑ์ต้นฉบับเป็นผ่าน'),
 ('NFR-05','คุณภาพ','ไม่เผยรหัส Wi-Fi/broker ใน source บน GitHub; NVS/ไฟล์ส่วนตัวถูก ignore; broker มี user/password\nเกณฑ์: repo ไม่มี credentials ลับ',True,'แล้ว',False,'ผ่านบางส่วน','broker ยืนยันตัวตน; ZIP source 447 ไฟล์ผ่านตรวจค่าความลับที่รู้จัก และไม่รวม private configs/firmware\nผลนี้เฉพาะชุด ZIP ยังไม่ได้ตรวจ repository ที่เผยแพร่จริง จึงไม่ติ๊กผ่านทั้งข้อ'),
 ('NFR-06','คุณภาพ','อุณหภูมิ ±3°C เทียบเครื่องอ้างอิง 10 จุด; BPM/SpO₂ ระบุว่าเบื้องต้น ไม่ใช่อุปกรณ์แพทย์',False,'บางส่วน',False,'รออ้างอิง','จอมีข้อความค่าประมาณ/ไม่ใช่เครื่องมือแพทย์ แต่ไม่มีเทอร์โมมิเตอร์ทดสอบ 10 จุด\nGarmin พักรอชาร์จ ใช้เทียบแบบผู้บริโภคได้ ไม่แทนเครื่องวัดอุณหภูมิ'),
 ('NFR-07','คุณภาพ','ฮาร์ดแวร์ต่อชุด ≤3,000 บาท\nเกณฑ์: รวม BOM ราคาจริง ครอบคลุมจอ INA226 และ QMI8658',False,'รอราคา',False,'รอรวมงบ','มีตาราง BOM แต่ราคายังว่างตามที่ผู้ใช้ขอเว้นไว้\nต้องเติมราคาซื้อจริงและชี้แจงขอบเขตชุดอุปกรณ์ก่อนตรวจงบ ไม่สร้างยอดประมาณเป็นผลผ่าน'),
 ('NFR-08','คุณภาพ','แยก Driver/Model/Service/UI; A/B rollback; OTA ห้ามเมื่อแบต <30%\nเกณฑ์: ตัดไฟขณะ OTA 3 ครั้ง แล้ว rollback/บูตกลับได้',True,'แล้ว',False,'รอตรวจ','มีโมดูล แบ่ง app0/app1 และเกณฑ์แบต 30% ในโค้ด\nยังไม่มีผลตัดไฟ OTA 3 ครั้งจริง การตรวจโค้ดหรือ host tests ไม่แทนผล rollback บนเครื่อง'),
]
def fmt(ms):
    minutes=int(ms//60000);return f'{minutes//60} h {minutes%60} min'
def para(container,text,style=None,size=None,bold=None):
    p=container.add_paragraph(style=style) if style else container.add_paragraph()
    r=p.add_run(text)
    if bold is not None:r.bold=bold
    previous=p._p.getprevious()
    if previous is not None and previous.tag==qn('w:tbl'):p.paragraph_format.space_before=Pt(6)
    if size:r.font.size=Pt(size)
    return p
def checkbox(cell,key,checked,caption):
    p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3)
    sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
    alias=OxmlElement('w:alias');alias.set(qn('w:val'),key);pr.append(alias)
    tag=OxmlElement('w:tag');tag.set(qn('w:val'),key);pr.append(tag)
    cb=etree_element('checkbox');v=etree_element('checked');v.set('{'+W14+'}val','1' if checked else '0');cb.append(v)
    for name,value in [('checkedState','2611'),('uncheckedState','2610')]:
        state=etree_element(name);state.set('{'+W14+'}val',value);state.set('{'+W14+'}font','Segoe UI Symbol');cb.append(state)
    pr.append(cb);sdt.append(pr);content=OxmlElement('w:sdtContent')
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Segoe UI Symbol')
    rp.append(fonts);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'30');rp.append(sz);r.append(rp)
    text=OxmlElement('w:t');text.text='☑' if checked else '☐';r.append(text);content.append(r);sdt.append(content);p._p.append(sdt)
    p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
    p.add_run(caption).font.size=Pt(10)
def etree_element(name):
    from lxml import etree
    return etree.Element('{'+W14+'}'+name)
def make_table(doc,rows):
    table=doc.add_table(rows=1,cols=5);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=[.65,2.05,.75,.8,2.75]
    for c,w in zip(table.columns,widths):c.width=Inches(w)
    for c,w,label in zip(table.rows[0].cells,widths,['ข้อ','ข้อกำหนดและเกณฑ์ย่อ','ทำแล้ว','ตรวจผ่าน','หลักฐานและสิ่งที่ยังขาด']):
        c.width=Inches(w);c.text=label
        for r in c.paragraphs[0].runs:r.bold=True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for entry in rows:
        ident,priority,criteria,done,done_label,passed,pass_label,note=entry
        cells=table.add_row().cells
        cells[0].text=ident+'\n'+priority;cells[1].text=criteria;cells[4].text=note
        cells[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        checkbox(cells[2],ident+'_implemented',done,done_label);checkbox(cells[3],ident+'_accepted',passed,pass_label)
    for idx,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,(cell,width) in enumerate(zip(row.cells,widths)):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            prop=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if idx==0 else 'FFFFFF');prop.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'85' if side in ['top','bottom'] else '80');v.set(qn('w:type'),'dxa');margins.append(v)
            prop.append(margins);borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side)
                for name,value in {'val':'single','sz':'4','color':'D9D9D9'}.items():v.set(qn('w:'+name),value)
                borders.append(v)
            prop.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
                for r in p.runs:
                    if r.font.name!='Segoe UI Symbol' and not r._r.xpath('.//w:sdt'):r.font.size=Pt(11 if j in [1,4] else 10)
    return table
def main():
    OUT.mkdir(parents=True,exist_ok=True)
    freeze=OUT/'frozen_sources.json'
    if not freeze.exists():
        original=json.loads((ROOT/'output/deadline_20261018/requirements_original.json').read_text(encoding='utf8'))
        source=Path(original['source']);actual=Document(source)
        actual_rows=[[c.text.strip() for c in row.cells] for t in actual.tables for row in t.rows]
        for req in original['requirements']:assert any(r[0]==req['id'] and req['requirement'] in r and req['acceptance'] in r for r in actual_rows),req['id']
        files=['output/accumulated_soak/status.json','output/deadline_20261018/source_delivery/manifest.json','output/deadline_20261018/sd_verified/env_1791956522_e47806b1.csv.audit.json','output/deadline_20261018/sd_verified/excel_open.json','output/deadline_20261018/hub_sd_user_confirmation.json','output/deadline_20261018/battery_test_deferred.json','output/deadline_20261018/sensor_comparison_baseline_summary.json']
        snapshots={name:json.loads((ROOT/name).read_text(encoding='utf-8-sig')) for name in files}
        obj={'created_utc':datetime.now(timezone.utc).isoformat(),'original_source':str(source),'original_source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'original_requirements':original['requirements'],'files_sha256':{name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in files},'snapshots':snapshots}
        freeze.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf8')
    frozen=json.loads(freeze.read_text(encoding='utf8'));soak=frozen['snapshots']['output/accumulated_soak/status.json']
    when=datetime.fromisoformat(soak['updated_utc'].replace('Z','+00:00')).astimezone(timezone(timedelta(hours=7)))
    rows=[]
    for row in ROWS:
        row=list(row)
        if row[-1]=='SOAK_DYNAMIC':row[-1]=f"MQTT สดสะสม {fmt(soak['credited_ms'])}; ต่อเนื่องยาวที่สุด {fmt(soak['longest_continuous_ms'])}; ช่วงปัจจุบัน {fmt(soak['current_continuous_ms'])}\nยังไม่ครบ 24 h ต่อเนื่อง และจอใหญ่เพิ่งรีบูตเพื่อลง SD ต้องแยกขอบเขตจากทั้งระบบและแบต"
        rows.append(row)
    assert len(rows)==16 and {r[0] for r in rows}=={r['id'] for r in frozen['original_requirements']}
    doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=section.right_margin=Inches(.75);section.top_margin=section.bottom_margin=Inches(.65)
    section.header_distance=section.footer_distance=Inches(.3)
    for style in ['Normal','Title','Heading 1']:
        s=doc.styles[style];s.font.name=FONT;s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);s.element.get_or_add_rPr().rFonts.set(qn('w:cs'),FONT)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11.5);doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    header=section.header.paragraphs[0];header.text='KU Watch   ตารางตรวจรับตาม Requirements V2';header.runs[0].font.size=Pt(9)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f'14 ตุลาคม 2569 {when:%H:%M} น.   หน้า ').font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='ตารางเทียบความต้องการและผลตรวจรับ KU Watch';doc.core_properties.author='KU Watch'
    para(doc,'ตารางเทียบความต้องการ KU Watch',style='Title')
    para(doc,f'สถานะ 14 ตุลาคม 2569 เวลา {when:%H:%M} น.   WATCH 2.0.8   HUB 1.0.7',size=10.5)
    para(doc,'เทียบข้อ FR และ NFR ทั้ง 16 ข้อกับ Requirements Plan V2 เดิม ตารางย่อข้อความเพื่ออ่านง่ายโดยคงรหัสและเกณฑ์ตัวเลข ต้นแบบทำงานได้หลายส่วน แต่ยังมีผลตรวจรับที่ไม่ครบ')
    para(doc,'ช่องทำแล้ว ☑ หมายถึงส่วนระบบพร้อมใช้ตามหมายเหตุ ส่วนช่องตรวจผ่าน ☑ ต้องมีหลักฐานถึงเกณฑ์ที่ระบุ ช่องว่าง ☐ ไม่ถือว่าผ่าน; ช่องติ๊กแก้ไขได้ใน Word',size=10.5)
    para(doc,'ข้อกำหนดหลัก FR 01 ถึง FR 04',style='Heading 1');make_table(doc,rows[:4])
    para(doc,'การเปลี่ยนอุปกรณ์: ผู้ใช้ยืนยันว่าผู้สอนยอมรับแฟลชบนตัวเล็ก + MicroSD บนจอใหญ่ และเซ็นเซอร์ที่ปรับใช้ การยอมรับนี้ไม่แทนผลทดสอบ ตัวเล็กมี GY-51; QMI8658/PCF85063 อยู่บนจอใหญ่',size=10)
    doc.add_page_break();para(doc,'ข้อกำหนดเพิ่มเติม FR 05 ถึง FR 08',style='Heading 1');make_table(doc,rows[4:8])
    para(doc,'Must/Should/Could ยึดตารางจัดลำดับในเอกสารเดิม FR-08 เป็น Could แต่ข้อ rollback ใน NFR-08 ยังต้องมีผลตรวจรับถ้าจะสรุปว่าครบทุกข้อ',size=10.5)
    doc.add_page_break();para(doc,'คุณภาพระบบ NFR 01 ถึง NFR 04',style='Heading 1');make_table(doc,rows[8:12])
    para(doc,'การเก็บ MQTT จาก USB ตรวจการรับข้อมูลและความเสถียรในขอบเขตที่บันทึก ไม่พิสูจน์อายุแบตเตอรี่ และไม่รวมหลายช่วงเป็น 24 ชั่วโมงต่อเนื่อง รอบเดิมที่ไม่ผ่านยังเก็บแยกไว้',size=10.5)
    doc.add_page_break();para(doc,'คุณภาพระบบ NFR 05 ถึง NFR 08',style='Heading 1');make_table(doc,rows[12:])
    para(doc,'หลักฐานและการใช้ตาราง',style='Heading 1')
    para(doc,'FR-04 ใช้ไฟล์จาก MicroSD จริงและผลเปิด Microsoft Excel; FR-02 มีคำยืนยันผู้ใช้/serial; รอบยาวใช้ status.json; ความปลอดภัยใช้ manifest ของ source ZIP; เกณฑ์ FR/NFR ตรวจตรงกับเอกสารต้นฉบับแล้ว',size=10.5)
    para(doc,'ต้นฉบับอ้างอิง Smart Environment Watch Requirements Plan V2 (1).docx ไม่แก้เกณฑ์เดิม ข้อพักทดสอบยังไม่ผ่าน การเทียบ Garmin รอชาร์จ อุณหภูมิยังไม่มีเครื่องอ้างอิง ราคาจริงและภาพ/วิดีโอยังต้องเติม',size=10.5)
    file=OUT/'KU_Watch_ตารางเช็กRequirements_14ตุลาคม2569.docx';doc.save(file)
    comparison={'snapshot_utc':soak['updated_utc'],'snapshot_thai':when.isoformat(),'boundary':'Checkboxes distinguish implemented system work from evidence-backed acceptance. Only FR-04 is checked as a short observation pass; partial and deferred tests are not passed.','requirements':[{'id':r[0],'priority':r[1],'criteria_summary':r[2],'implemented_checked':r[3],'implementation_label':r[4],'acceptance_checked':r[5],'acceptance_label':r[6],'evidence_and_remaining':r[7],'original_requirement':next(x['requirement'] for x in frozen['original_requirements'] if x['id']==r[0]),'original_acceptance':next(x['acceptance'] for x in frozen['original_requirements'] if x['id']==r[0])} for r in rows]}
    (OUT/'comparison.json').write_text(json.dumps(comparison,ensure_ascii=False,indent=2),encoding='utf8')
    (OUT/'manifest.json').write_text(json.dumps({'file':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'requirements':16,'editable_checkboxes':32,'original_source_sha256':frozen['original_source_sha256']},indent=2),encoding='utf8')
    print(file)
if __name__=='__main__':main()
